import importlib.util
from pathlib import Path
import tempfile
import unittest
import openpyxl
spec=importlib.util.spec_from_file_location('history',Path(__file__).with_name('preparar-historico-financeiro.py'))
history=importlib.util.module_from_spec(spec);spec.loader.exec_module(history)
class FinancialConverterTests(unittest.TestCase):
    def test_explicit_summary_reconciles_totals_without_reimporting_movements(self):
        import datetime
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'teste.xlsx';wb=openpyxl.Workbook();s=wb.active;s.title='Planilha1'
            s.append(['CODIGO DA TRANSACAO','DATA','TIPO','DESCRICAO','VALOR','Dia da Semana','Comissão?','Tipo_01','Tipo_02','Tipo_03'])
            s.append(['id-1',datetime.datetime(2026,9,1),'Vendas','Crédito',900,'Terça','Não','Receita','Bar','Bar'])
            s.append(['externa',datetime.datetime(2026,9,30),'Compra','Fornecedor',-800,'Quarta','Não','Despesa','Compras','Compras'])
            s.append(['externa',datetime.datetime(2026,9,30),'Compra','Outro',-200,'Quarta','Não','Despesa','Compras','Compras'])
            summary=wb.create_sheet('Planilha2')
            for row,label,value in [(31,'Despesa',-1000),(48,'Receita',900),(51,'Total Geral',-100)]:summary.cell(row,3,label);summary.cell(row,4,value)
            wb.save(path);result=history.prepare_statement_summary(path)
            self.assertEqual(len(result['records']),2)
            self.assertEqual({r['data']['series'] for r in result['records']},{'consolidado'})
            self.assertEqual(sum(r['data']['amountCents'] for r in result['records']),-10000)
            self.assertEqual(result['records'][0]['data']['metric'],'receitas-consolidadas')
            self.assertEqual(result['records'][0]['data']['sourceCell'],'Planilha2!D48')
            self.assertEqual(result['records'][0]['data']['periodEnd'],'2026-09-30')
            # Repeated external identifiers cannot become bank transactions through this mode.
            with self.assertRaises(ValueError):history.prepare(path)
            summary['D31']=-999;wb.save(path)
            with self.assertRaises(ValueError):history.prepare_statement_summary(path)
            summary['D31']=-1000;summary['D51']=None;wb.save(path)
            with self.assertRaises(ValueError):history.prepare_statement_summary(path)
            summary['D51']=-100;s['B4']=datetime.datetime(2026,10,1);wb.save(path)
            with self.assertRaises(ValueError):history.prepare_statement_summary(path)
    def test_bank_transactions_keep_dates_and_signs_exclude_pivot_and_reject_duplicates(self):
        import datetime
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'teste.xlsx';wb=openpyxl.Workbook();s=wb.active;s.title='Planilha1'
            s.append(['CODIGO DA TRANSACAO','DATA','TIPO','DESCRICAO','VALOR','Dia da Semana','Comissão?','Tipo_01','Tipo_02','Tipo_03'])
            s.append(['id-1',datetime.datetime(2026,9,1),'Vendas','Crédito',123.45,'Terça','Não','Receita','Bar','Bar'])
            s.append(['id-2',datetime.datetime(2026,9,3),'Compra','Fornecedor',-10,'Quinta','Não','Despesa','Compras','Compras'])
            wb.create_sheet('Planilha2').append(['Total',9999]);wb.save(path)
            result=history.prepare(path);self.assertEqual(len(result['records']),2)
            self.assertEqual(result['records'][1]['data']['amountCents'],-1000)
            self.assertEqual(result['records'][1]['data']['periodStart'],'2026-09-03')
            self.assertEqual(result['records'][1]['data']['series'],'pagbank-conta')
            s.append(list(s.values)[1]);wb.save(path)
            with self.assertRaises(ValueError):history.prepare(path)
    def test_preserves_cents_and_rejects_invalid_precision(self):
        self.assertEqual(history.cents(12.34),1234)
        self.assertEqual(history.cents(-12.34),-1234)
        for value in ['NaN','Infinity',1.123]:
            with self.assertRaises(ValueError):history.cents(value)
    def test_rejects_unrelated_business_spreadsheet(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'outro.xlsx';openpyxl.Workbook().save(path)
            with self.assertRaises(ValueError):history.prepare(path)
    def test_does_not_turn_unpaid_months_into_receipts_or_import_month_totals_as_classes(self):
        import datetime
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'teste.xlsx';wb=openpyxl.Workbook();wb.remove(wb.active)
            s=wb.create_sheet('Controle Alunos');s.append(['Mês','Nome','Qtd','Dias','Turma','Turma','Valor Pago','Status']);s.append([202608,'Aluno teste',1,None,None,None,123.45,'Não Pago'])
            s=wb.create_sheet('Controle Mensalistas');s.append(['Mês','Dia','Hora','Nome','Valor Pago','Status']);s.append([202608,None,None,'Mensalista teste',100,'Pago'])
            s=wb.create_sheet('Controle Aulas');s.append(['Data','Dia','Hora','Qtd','Valor Escalonável','Status']);s.append([datetime.datetime(2026,8,3),'segunda','18:00',2,10,'Pago']);s.append([202608,None,None,None,10,'Total'])
            wb.save(path);result=history.prepare(path)
            self.assertEqual(len(result['records']),3)
            first=result['records'][0]['data'];self.assertEqual(first['state'],'Não Pago');self.assertEqual(first['amountCents'],12345)
            self.assertEqual(first['periodEnd'],'2026-08-31');self.assertEqual(first['sourceCell'],'Controle Alunos!G2')
            self.assertEqual(result['sourceSha256'],history.hashlib.sha256(path.read_bytes()).hexdigest())
if __name__=='__main__':unittest.main()
