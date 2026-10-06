#!/usr/bin/env python3
"""Leitura de controles XLSX da arena; valores históricos, sem gerar cobranças ou movimentos de estoque.

Requer openpyxl somente na estação de preparação. Fontes nunca são alteradas.
Resumos, detalhes, estimativas e saldos são séries independentes, sem soma entre elas.
"""
import argparse
import calendar
import datetime as dt
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import openpyxl


def cents(value):
    if value is None or value == '': return None
    amount = Decimal(str(value))
    if not amount.is_finite():raise ValueError('Valor financeiro não finito.')
    rounded = amount.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    if not amount.is_finite() or abs(amount-rounded) > Decimal('0.000001') or abs(amount)>999999999:
        raise ValueError('Valor financeiro inválido ou precisão inesperada.')
    return int(rounded*100)


def month(value):
    text = str(value)
    if len(text)!=6 or not text.isdigit(): raise ValueError('Competência inválida.')
    return dt.date(int(text[:4]),int(text[4:]),1)


def serial(value):
    return value.isoformat() if isinstance(value,(dt.date,dt.datetime,dt.time)) else value


def prepare(source):
    wb=openpyxl.load_workbook(source,read_only=True,data_only=True)
    records=[]
    def add(sheet,row,col,series,metric,label,state,period,value,grain='month',notes=''):
        amount=cents(value)
        if amount is None:return
        end=period.replace(day=calendar.monthrange(period.year,period.month)[1]) if grain in ('month','estimate') else period
        cell=f'{sheet}!{openpyxl.utils.get_column_letter(col)}{row}'
        data={'entity':'financial-observation-v1','currency':'BRL','sourceCell':cell,'series':series,'metric':metric,
              'label':str(label).strip(),'state':str(state),'periodStart':period.isoformat(),'periodEnd':end.isoformat(),
              'grain':grain,'amountCents':amount,'notes':str(notes)[:2000]}
        records.append({'sheetName':sheet+' • '+openpyxl.utils.get_column_letter(col),'rowNumber':row,
                        'recordType':'reference-data','externalId':'finance:'+cell,'data':data})
    if {'Controle Aulas','Controle Alunos','Controle Mensalistas'}.issubset(wb.sheetnames):
        for sheet,series,amountcol,namecol,statuscol in [('Controle Alunos','alunos',7,2,8),('Controle Mensalistas','mensalistas',5,4,6)]:
            s=wb[sheet]
            expected='Valor Pago'
            if s.cell(1,amountcol).value!=expected:raise ValueError('Cabeçalho de controle não reconhecido.')
            for row,values in enumerate(s.values,1):
                if row==1 or not values[0] or not values[namecol-1]:continue
                first=month(values[0]);state=values[statuscol-1] or 'Não informado'
                notes=json.dumps({str(s.cell(1,i+1).value or f'Coluna {i+1}'):serial(v) for i,v in enumerate(values) if v is not None},ensure_ascii=False)
                add(sheet,row,amountcol,series,'valor-informado',values[namecol-1],state,first,values[amountcol-1],notes=notes)
        s=wb['Controle Aulas']
        if s.cell(1,5).value!='Valor Escalonável':raise ValueError('Cabeçalho das aulas não reconhecido.')
        for row,values in enumerate(s.values,1):
            if row==1 or not isinstance(values[0],dt.datetime):continue
            add(s.title,row,5,'aulas','valor-escalonavel',f'Aula {values[2]}',values[5] or 'Não informado',values[0].date(),values[4],grain='day',
                notes=json.dumps({'quantidadeAlunos':values[3],'dia':values[1]},ensure_ascii=False))
    elif 'Planilha1' in wb.sheetnames and [wb['Planilha1'].cell(1,c).value for c in range(1,11)] == ['CODIGO DA TRANSACAO','DATA','TIPO','DESCRICAO','VALOR','Dia da Semana','Comissão?','Tipo_01','Tipo_02','Tipo_03']:
        seen=set()
        for row,values in enumerate(wb['Planilha1'].values,1):
            if row==1 or not any(v is not None for v in values):continue
            code,date,kind,description,value=values[:5]
            if not code or code in seen or not isinstance(date,dt.datetime) or value is None or values[7] not in ('Receita','Despesa'):
                raise ValueError('Movimento bancário sem identidade, data, valor, classificação ou com duplicidade.')
            amount=cents(value)
            if amount==0 or (amount>0)!=(values[7]=='Receita'):raise ValueError('Sinal e classificação bancária divergentes.')
            seen.add(code)
            add('Planilha1',row,5,'pagbank-conta','entradas-extrato' if amount>0 else 'saidas-extrato',description,'Informado no extrato',date.date(),value,grain='day',notes=json.dumps({'transactionCode':code,'tipo':kind,'categoria':values[8],'classificacao':values[9],'comissao':values[6]},ensure_ascii=False))
        # Pivot subtotals are intentionally excluded: they can include off-account expenses.
        # Such values require explicit reconciliation and an independently reviewed staging batch.
    elif '202608' in wb.sheetnames:
        s=wb['202608']
        if [s.cell(4,c).value for c in (6,7,8,9)]!=['Descrição','Valor de 202606','Valor Pago 202607','Valor Pago 202608']:
            raise ValueError('Cabeçalho do consolidado não reconhecido; reveja o mapeamento.')
        for row in range(5,28):
            label=s.cell(row,6).value;category=s.cell(row,11).value
            if not label or not category:raise ValueError('Linha do consolidado sem classificação.')
            metric='vendas-bar-bruto' if category=='Receita Bar Bruta' else 'despesas' if category.startswith('Despesa') else 'receitas-arena'
            for col,reference in [(7,202606),(8,202607),(9,202608)]:
                add(s.title,row,col,'consolidado',metric,label,'Referência' if col==7 else 'Informado na planilha',month(reference),s.cell(row,col).value,
                    grain='estimate' if col==7 else 'month',notes=f'{category}; competência mensal; não confirma liquidação bancária.')
        # Segundo quadro contém parcelas adicionais. É uma visão alternativa, jamais somada ao primeiro.
        for row in range(70,97):
            label=s.cell(row,6).value;category=s.cell(row,8).value
            if not label or not category:continue
            value=s.cell(row,7).value
            metric='vendas-bar-bruto' if str(label).strip()=='Vendas Bar Bruto' else 'despesas' if category.startswith('Despesa') else 'receitas-arena'
            add(s.title,row,7,'consolidado-com-dividas',metric,label,'Informado na planilha',month(202608),value,
                notes=json.dumps({'tipo':category,'vencimento':serial(s.cell(row,9).value),'pagamento':serial(s.cell(row,10).value)},ensure_ascii=False))
        for row in range(36,48):
            if isinstance(s.cell(row,11).value,dt.datetime):
                add(s.title,row,13,'compras-detalhadas',str(s.cell(row,14).value),s.cell(row,12).value,'Informado na planilha',s.cell(row,11).value.date(),s.cell(row,13).value,grain='day',notes=s.cell(row,15).value or '')
            if isinstance(s.cell(row,6).value,dt.datetime):
                add(s.title,row,8,'servicos-detalhados',str(s.cell(row,7).value),s.cell(row,7).value,'Informado na planilha',s.cell(row,6).value.date(),s.cell(row,8).value,grain='day',notes=s.cell(row,9).value or '')
        for row in range(53,60):
            label=s.cell(row,7).value
            add(s.title,row,8,'saldos',str(label),label,'Fotografia da planilha',s.cell(row,6).value.date(),s.cell(row,8).value,grain='snapshot',notes=s.cell(row,9).value or '')
    else:raise ValueError('Arquivo não reconhecido como controle financeiro da arena.')
    if not 1<=len(records)<=2000:raise ValueError('Lote fora do limite.')
    if len({r['data']['sourceCell']for r in records})!=len(records):raise ValueError('Célula duplicada.')
    return {'sourceName':source.name+'-financeiro-v1.json','sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest(),'records':records}


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('source',type=Path);parser.add_argument('output',type=Path)
    args=parser.parse_args();package=prepare(args.source)
    args.output.write_text(json.dumps(package,ensure_ascii=False,indent=2)+'\n');args.output.chmod(0o600)
    print(json.dumps({'records':len(package['records']),'sha256':package['sourceSha256']}))
